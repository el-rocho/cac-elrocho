"""Pruebas de exportación a Microsoft Excel (.xlsx) con filtro de fechas y gráficos nativos."""

from __future__ import annotations

from io import BytesIO
import unittest
from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.models import Base, Informe
from app.seed_data import run_seed
from app.services.excel_export_service import generate_clinical_excel_bytes


PROJECT_ROOT = Path(__file__).resolve().parent.parent


class TestExcelExport(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(bind=self.engine)
        self.SessionLocal = sessionmaker(bind=self.engine)
        self.db = self.SessionLocal()
        run_seed(self.db)

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def test_genera_excel_completo_con_tres_hojas_y_graficos_nativos(self):
        excel_bytes, eff_start, eff_end = generate_clinical_excel_bytes(self.db)
        self.assertGreater(len(excel_bytes), 5000)

        wb = load_workbook(BytesIO(excel_bytes))
        self.assertEqual(wb.sheetnames, ["Resumen", "Datos", "Gráficos"])

        ws_resumen = wb["Resumen"]
        resumen_values = [
            str(cell.value)
            for row in ws_resumen.iter_rows()
            for cell in row
            if cell.value is not None
        ]
        resumen_text = "\n".join(resumen_values)

        for titulo_kpi in (
            "METABOLISMO GLUCÍDICO",
            "PERFIL LIPÍDICO",
            "FUNCIÓN RENAL",
            "FUNCIÓN HEPÁTICA",
            "HEMOGRAMA / HIERRO",
            "FUNCIÓN TIROIDEA",
        ):
            self.assertIn(titulo_kpi, resumen_text)
        self.assertIn("OTROS VALORES DE INTERÉS CLÍNICO", resumen_text)

        ws_datos = wb["Datos"]
        self.assertEqual(ws_datos.freeze_panes, "D4")
        self.assertEqual(ws_datos["A3"].value, "Parámetro")
        self.assertEqual(ws_datos["B3"].value, "Unidad de medida")
        self.assertEqual(ws_datos["C3"].value, "Referencia oficial")

        ws_graficos = wb["Gráficos"]
        self.assertGreater(len(ws_graficos._charts), 0)

    def test_filtra_correctamente_por_rango_de_fechas(self):
        informes = self.db.query(Informe).order_by(Informe.fecha.asc()).all()
        self.assertGreaterEqual(len(informes), 5)
        start_f = informes[1].fecha
        end_f = informes[3].fecha

        excel_bytes, eff_start, eff_end = generate_clinical_excel_bytes(
            self.db, fecha_inicio=start_f, fecha_fin=end_f
        )
        self.assertEqual(eff_start, start_f)
        self.assertEqual(eff_end, end_f)

        wb = load_workbook(BytesIO(excel_bytes))
        ws_datos = wb["Datos"]
        # 3 columnas fijas + 3 fechas filtradas + 1 columna Promedio 18 meses = 7 columnas con encabezado en fila 3
        headers_row_3 = [
            ws_datos.cell(row=3, column=col).value
            for col in range(1, 15)
            if ws_datos.cell(row=3, column=col).value is not None
        ]
        self.assertEqual(len(headers_row_3), 3 + 3 + 1)
        self.assertEqual(headers_row_3[-1], "Promedio 18 meses")

    def test_rango_invalido_o_vacio_lanza_value_error(self):
        with self.assertRaises(ValueError):
            generate_clinical_excel_bytes(
                self.db, fecha_inicio="2026-12-31", fecha_fin="2025-01-01"
            )
        with self.assertRaises(ValueError):
            generate_clinical_excel_bytes(
                self.db, fecha_inicio="1990-01-01", fecha_fin="1991-01-01"
            )

    def test_interfaz_configuracion_incluye_tercera_tarjeta_excel_en_misma_fila(self):
        index = (PROJECT_ROOT / "app" / "static" / "index.html").read_text(encoding="utf-8")
        app_js = (PROJECT_ROOT / "app" / "static" / "app.js").read_text(encoding="utf-8")

        self.assertIn("lg:grid-cols-3", index)
        self.assertIn("Exportar informe Excel", index)
        for elem_id in ("excelStartDate", "excelEndDate", "btnExportExcel"):
            self.assertIn(f'id="{elem_id}"', index)

        for fn_name in ("function downloadExcelReport", "function syncExcelDateBounds", "function updateExcelDateRangeInfo"):
            self.assertIn(fn_name, app_js)


if __name__ == "__main__":
    unittest.main()
